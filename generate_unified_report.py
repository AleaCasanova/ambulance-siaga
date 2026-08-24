import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
wb.remove(wb.active) # remove default sheet

# Data: Fitur Terpadu / Alur Kerja Realistis (Tidak lompat-lompat FE/BE per hari)
weeks_data = [
    {
        "week_num": "1",
        "title": "Minggu ke : 1 | Hari : Senin – Jumat | Tanggal : 3 Agustus – 7 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Setup Lingkungan & Arsitektur Basis Data Dasar",
             "Menginisialisasi proyek Laravel 11, konfigurasi database MySQL, pembuatan skema migrasi tabel dasar (users, roles, masyarakat, supir, ambulans, rumah_sakit), dan penyiapan layout antarmuka aplikasi.",
             "Screenshot struktur folder Laravel & daftar migrasi database",
             "Selesai (Environment & Database siap)"),
            (2, "Pengembangan Modul Pemesanan Pasien & Integrasi Peta Leaflet",
             "Membangun formulir pemesanan darurat (order-create.blade.php), integrasi peta interaktif Leaflet.js dengan marker penjemputan draggable, dan controller penyimpanan data pesanan ke database.",
             "Screenshot form pemesanan dan interaksi peta penjemputan",
             "Selesai (Formulir & peta jemput berfungsi)"),
            (3, "Pengembangan Modul Autentikasi Multi-Role & Verifikasi Foto",
             "Membuat antarmuka dan backend login/register multi-role (Masyarakat, Supir, Mitra), middleware hak akses, serta penanganan unggah file foto kondisi darurat pasien.",
             "Screenshot halaman login/register & upload berkas",
             "Selesai (Autentikasi multi-role aktif)"),
            (4, "Pengembangan Sistem Verifikasi Akun (OTP Email)",
             "Membangun antarmuka verifikasi OTP (verify-otp.blade.php), logika backend token OTP acak pada tabel otp_verifications, serta pengiriman email notifikasi via SMTP Laravel Mail.",
             "Screenshot form verifikasi OTP & template email OTP",
             "Selesai (Verifikasi akun OTP berfungsi)"),
            (5, "Pengembangan Modul Donasi Publik & Sinkronisasi Repositori",
             "Membangun antarmuka dan backend halaman donasi publik (donasi.blade.php), pengujian alur pemesanan darurat tahap 1, serta resolusi konflik git merge tim.",
             "Screenshot halaman donasi & riwayat commit git",
             "Selesai (Modul donasi aktif & repo sinkron)")
        ]
    },
    {
        "week_num": "2",
        "title": "Minggu ke : 2 | Hari : Senin – Jumat | Tanggal : 10 Agustus – 14 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Penyempurnaan Navigasi, Notifikasi Lonceng & Hak Akses",
             "Merancang ulang navbar aplikasi dengan icon lonceng notifikasi dinamis, perbaikan judul header, serta refaktorisasi 4 role utama (Admin, Operator, Supir, Masyarakat) pada database & middleware.",
             "Screenshot navbar notifikasi & menu navigasi per role",
             "Selesai (Navigasi & role terstruktur)"),
            (2, "Desain Ulang Hero Section & Tampilan Beranda Utama",
             "Mengembangkan tampilan modern dan responsif pada halaman Beranda (Hero section interaktif, kartu keunggulan armada darurat, dan alur panggilan siaga).",
             "Screenshot tampilan baru Landing Page Beranda",
             "Selesai (Beranda responsif & modern)"),
            (3, "Pengembangan Landing Page Portal Supir & Mitra Lembaga",
             "Membangun halaman informasi publik dan registrasi terpisah untuk supir ambulans (register-supir.blade.php) serta mitra lembaga (register-mitra.blade.php).",
             "Screenshot portal supir dan portal mitra",
             "Selesai (Portal supir & mitra siap)"),
            (4, "Perancangan Skema Database Donasi & Konfigurasi API Midtrans",
             "Membuat migrasi tabel donasis, konfigurasi environment key Midtrans Snap API, dan penyiapan webhook handler status transaksi donasi.",
             "Screenshot struktur tabel donasis & konfigurasi API",
             "Selesai (Database & konfigurasi API siap)"),
            (5, "Integrasi Payment Gateway Midtrans Snap & Manajemen Mitra",
             "Mengimplementasikan integrasi popup Midtrans Snap (QRIS, VA, E-Wallet), migrasi tabel mitras, role Mitra lembaga, dan status alur 'menunggu_konfirmasi_supir'.",
             "Screenshot modal Midtrans Snap payment & dashboard mitra",
             "Selesai (Transaksi donasi online berhasil)")
        ]
    },
    {
        "week_num": "3",
        "title": "Minggu ke : 3 | Hari : Senin – Jumat | Tanggal : 17 Agustus – 21 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Pengembangan Console Supir & Bottom Navigation Mobile",
             "Membangun dashboard kerja khusus supir (tugas-index.blade.php) yang dioptimasi untuk layar smartphone dengan mobile bottom navigation bar.",
             "Screenshot dashboard supir versi smartphone",
             "Selesai (Mobile console supir aktif)"),
            (2, "Pengembangan Fitur Triage Prioritas Darurat & Berkas Supir",
             "Menambahkan opsi prioritas darurat (Gawat Darurat / Normal) pada form pemesanan, validasi upload file SIM & STNK supir, serta pagination daftar data.",
             "Screenshot form upload SIM/STNK & badge prioritas order",
             "Selesai (Prioritas order & verifikasi supir aktif)"),
            (3, "Setup REST API Token (Sanctum) & Konsol Dispatch Operator",
             "Konfigurasi Laravel Sanctum untuk endpoint API autentikasi perangkat GPS dan penguatan logika dispatching penugasan armada ambulans terdekat.",
             "Screenshot konsol dispatch operator & route API",
             "Selesai (API token & dispatching siap)"),
            (4, "Pengembangan Live Tracking GPS Ambulans & Rute OSRM",
             "Membangun halaman live tracking pasien (order-tracking.blade.php) dengan visualisasi rute polyline jalan, pergerakan marker ambulans realtime, dan estimasi waktu tiba (ETA).",
             "Screenshot halaman live tracking pasien & rute jalan",
             "Selesai (Live tracking peta berfungsi)"),
            (5, "Penyempurnaan Modul Ulasan (Rating) & Kwitansi Donasi",
             "Mengembangkan modul rating bintang untuk ulasan masyarakat terhadap layanan supir/ambulans, halaman kwitansi bukti donasi sukses, dan detail riwayat pesanan.",
             "Screenshot form rating bintang & kwitansi donasi",
             "Selesai (Modul rating & kwitansi selesai)")
        ]
    },
    {
        "week_num": "4",
        "title": "Minggu ke : 4 | Hari : Senin – Jumat | Tanggal : 24 Agustus – 28 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Penyeragaman & Optimalisasi Peta Leaflet di Seluruh Role",
             "Menyeragamkan tile layer OpenStreetMap di seluruh role (Admin, Operator, Mitra, Supir, Pasien), mengaktifkan interaktivitas zoom scroll mouse pada Admin, dan memperbarui pin marker 3D status armada.",
             "Screenshot tampilan peta seragam di Admin, Operator, & Pasien",
             "Selesai (Peta seragam, bersih, & interaktif)"),
            (2, "Pengujian Terpadu Alur Pemesanan & Navigasi Ambulans End-to-End",
             "Melakukan uji coba komprehensif mulai dari pemesanan darurat pasien, notifikasi operator, penugasan armada, navigasi supir menuju pasien, hingga update status selesai.",
             "Screenshot alur status perjalanan ambulans dari awal sampai selesai",
             "Selesai (Alur end-to-end terverifikasi)")
        ]
    }
]

# Styles
title_font = Font(name='Segoe UI', size=14, bold=True, color='0F2742')
subtitle_font = Font(name='Segoe UI', size=11, bold=True, color='1F497D')
header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
cell_font = Font(name='Segoe UI', size=10, color='1E293B')
status_font = Font(name='Segoe UI', size=10, bold=True, color='166534')

header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
status_fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')

thin_border_side = Side(border_style='thin', color='CBD5E1')
cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

header_border_side = Side(border_style='medium', color='1F497D')
header_border = Border(left=thin_border_side, right=thin_border_side, top=header_border_side, bottom=header_border_side)

for w in weeks_data:
    ws = wb.create_sheet(title=f"Minggu {w['week_num']}")
    ws.views.sheetView[0].showGridLines = True

    # Title
    ws.cell(row=1, column=1, value="Form Detail Pekerjaan").font = title_font
    ws.cell(row=2, column=1, value=w['title']).font = subtitle_font

    # Headers exactly matching user template: No, Pekerjaan/Tugas, Uraian Pekerjaan, Gambar, Keterangan
    headers = ["No", "Pekerjaan/Tugas", "Uraian Pekerjaan", "Gambar", "Keterangan"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = header_border

    ws.row_dimensions[4].height = 28

    # Rows
    for row_idx, task in enumerate(w['tasks'], start=5):
        no, name, desc, img, status = task
        
        c1 = ws.cell(row=row_idx, column=1, value=no)
        c1.alignment = Alignment(horizontal='center', vertical='center')
        c1.font = cell_font
        c1.border = cell_border

        c2 = ws.cell(row=row_idx, column=2, value=name)
        c2.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        c2.font = Font(name='Segoe UI', size=10, bold=True, color='0F172A')
        c2.border = cell_border

        c3 = ws.cell(row=row_idx, column=3, value=desc)
        c3.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        c3.font = cell_font
        c3.border = cell_border

        c4 = ws.cell(row=row_idx, column=4, value=img)
        c4.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        c4.font = cell_font
        c4.border = cell_border

        c5 = ws.cell(row=row_idx, column=5, value=status)
        c5.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c5.font = status_font
        c5.fill = status_fill
        c5.border = cell_border

        ws.row_dimensions[row_idx].height = 55

    # Column widths
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 58
    ws.column_dimensions['D'].width = 38
    ws.column_dimensions['E'].width = 24

# Sheet 1: Khusus Rincian Pekerjaan Front-End
ws_fe = wb.create_sheet(title="Khusus Front-End", index=0)
ws_fe.views.sheetView[0].showGridLines = True
ws_fe.cell(row=1, column=1, value="DAFTAR PEKERJAAN FRONT-END (UI/UX & TAMPILAN)").font = title_font
ws_fe.cell(row=2, column=1, value="Proyek Sistem Informasi & Layanan Darurat Ambulans Siaga").font = subtitle_font

fe_headers = ["No", "Modul / Fitur", "Teknologi", "Uraian Pekerjaan Front-End", "Hasil"]
for col_idx, h in enumerate(fe_headers, start=1):
    cell = ws_fe.cell(row=4, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = header_border
ws_fe.row_dimensions[4].height = 28

fe_items = [
    (1, "Form Pemesanan Pasien & Peta", "Blade, Tailwind, Leaflet.js", "Membangun formulir pemesanan ambulans, integrasi picker peta Leaflet interaktif untuk penjemputan, dan modal verifikasi foto kondisi darurat.", "Selesai"),
    (2, "Landing Page & Beranda Utama", "Blade, Tailwind CSS, Alpine.js", "Merancang tampilan Hero section modern, alur darurat ambulans, kartu keunggulan armada, dan tata letak responsif mobile.", "Selesai"),
    (3, "Portal Supir & Mobile Navigation", "Blade, Tailwind, Alpine.js", "Mendesain console dashboard supir dengan mobile bottom navigation bar yang ramah sentuhan di layar smartphone.", "Selesai"),
    (4, "Live Tracking Peta Ambulans", "Leaflet.js, Alpine.js, CSS", "Visualisasi realtime posisi ambulans, garis rute jalan OSRM ke rumah sakit rujukan, dan badge estimasi waktu kedatangan (ETA).", "Selesai"),
    (5, "Modal Checkout Midtrans Snap", "Midtrans Snap JS, Tailwind", "Integrasi antarmuka popup checkout donasi multi-metode (QRIS, VA, E-Wallet) dan halaman kwitansi bukti transfer online.", "Selesai"),
    (6, "Dashboard Admin & Operator", "Blade, Tailwind, Leaflet.js", "Mendesain radar armada ambulans aktif, tabel riwayat pesanan, kartu ringkasan statistik, dan modal assignment penugasan supir.", "Selesai")
]

for r_idx, r_data in enumerate(fe_items, start=5):
    for c_idx, val in enumerate(r_data, start=1):
        cell = ws_fe.cell(row=r_idx, column=c_idx, value=val)
        cell.font = cell_font
        cell.border = cell_border
        if c_idx == 1:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx == 5:
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font = status_font
            cell.fill = status_fill
        else:
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    ws_fe.row_dimensions[r_idx].height = 42

ws_fe.column_dimensions['A'].width = 6
ws_fe.column_dimensions['B'].width = 28
ws_fe.column_dimensions['C'].width = 26
ws_fe.column_dimensions['D'].width = 62
ws_fe.column_dimensions['E'].width = 16

# Sheet 2: Khusus Rincian Pekerjaan Back-End
ws_be = wb.create_sheet(title="Khusus Back-End", index=1)
ws_be.views.sheetView[0].showGridLines = True
ws_be.cell(row=1, column=1, value="DAFTAR PEKERJAAN BACK-END (SERVER, DATABASE & API)").font = title_font
ws_be.cell(row=2, column=1, value="Proyek Sistem Informasi & Layanan Darurat Ambulans Siaga").font = subtitle_font

be_headers = ["No", "Modul / Fitur", "Teknologi", "Uraian Pekerjaan Back-End", "Hasil"]
for col_idx, h in enumerate(be_headers, start=1):
    cell = ws_be.cell(row=4, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = PatternFill(start_color='B45309', end_color='B45309', fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = header_border
ws_be.row_dimensions[4].height = 28

be_items = [
    (1, "Arsitektur Database & Migrasi", "MySQL, Laravel Migration, Eloquent", "Merancang skema database (users, roles, supir, ambulans, pemesanan, donasis, mitras), relasi tabel, dan indeks performa query.", "Selesai"),
    (2, "Sistem Autentikasi & OTP Email", "Laravel Auth, SMTP Mail, Session", "Membangun logika registrasi multi-role, pembuatan dan validasi kode token OTP 6-digit via email, serta middleware otorisasi.", "Selesai"),
    (3, "Payment Gateway & Webhook", "Midtrans Snap API, HTTP Client", "Menghubungkan API Midtrans dengan Server Key terenkripsi, pembuatan token transaksi, serta webhook callback status pembayaran.", "Selesai"),
    (4, "Dispatching & Alur Pemesanan", "Livewire Component, Service Layer", "Mengembangkan logika penugasan armada ambulans terdekat ke pasien dan pengelolaan alur status darurat (menunggu -> diproses -> selesai).", "Selesai"),
    (5, "REST API GPS & Laravel Sanctum", "Laravel Sanctum, JSON API", "Membuat endpoint API token untuk pengiriman koordinat GPS driver secara berkala dan integrasi layanan routing OSRM.", "Selesai"),
    (6, "Audit Trail & Rekapitulasi Data", "Laravel Eloquent, Event Listener", "Pencatatan riwayat aksi sistem (LogAktivitas) serta kalkulasi agregasi data statistik keuangan donasi dan performa supir.", "Selesai")
]

for r_idx, r_data in enumerate(be_items, start=5):
    for c_idx, val in enumerate(r_data, start=1):
        cell = ws_be.cell(row=r_idx, column=c_idx, value=val)
        cell.font = cell_font
        cell.border = cell_border
        if c_idx == 1:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx == 5:
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font = status_font
            cell.fill = status_fill
        else:
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    ws_be.row_dimensions[r_idx].height = 42

ws_be.column_dimensions['A'].width = 6
ws_be.column_dimensions['B'].width = 28
ws_be.column_dimensions['C'].width = 28
ws_be.column_dimensions['D'].width = 62
ws_be.column_dimensions['E'].width = 16

# Save files
for filename in ["laporan_mingguan_magang_v3.xlsx", "laporan_mingguan_magang_terpadu.xlsx"]:
    wb.save(filename)
    print(f"Saved {filename}")

try:
    wb.save("laporan_mingguan_magang.xlsx")
    print("Saved laporan_mingguan_magang.xlsx")
except Exception:
    pass
