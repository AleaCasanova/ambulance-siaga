import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
wb.remove(wb.active) # remove default sheet

# Data separated with Category: Front-End vs Back-End
weeks_data = [
    {
        "week_num": "1",
        "title": "Minggu ke : 1 | Hari : Senin – Jumat | Tanggal : 3 Agustus – 7 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Back-End", "Inisialisasi Proyek & Perancangan Basis Data Dasar",
             "Setup framework Laravel 11, konfigurasi database MySQL, pembuatan migrasi tabel dasar (users, roles, masyarakat, supir, ambulans, rumah_sakit), dan pemodelan Eloquent.",
             "Screenshot struktur migration database & konfigurasi .env",
             "Selesai (Database & core Laravel siap)"),
            (2, "Front-End", "Pembuatan Form Pemesanan Ambulans & Integrasi Peta",
             "Membangun antarmuka formulir pemesanan darurat (order-create.blade.php), styling Tailwind CSS, dan integrasi peta Leaflet.js dengan marker penjemputan draggable.",
             "Screenshot UI form pemesanan & interaksi peta Leaflet",
             "Selesai (Antarmuka pemesanan siap)"),
            (3, "Back-End", "Logika Autentikasi Multi-Role & Unggah Berkas Verifikasi",
             "Mengembangkan backend controller login/register multi-role, middleware otorisasi, dan validasi penyimpanan file foto kondisi gawat darurat pasien.",
             "Screenshot controller auth & alur penyimpanan file foto",
             "Selesai (Autentikasi & upload foto aktif)"),
            (4, "Back-End", "Sistem Verifikasi Akun Berbasis OTP / Email",
             "Implementasi logika pembuatan OTP acak, migrasi tabel otp_verifications, pengiriman email via SMTP Laravel Mail, dan validasi masa aktif token OTP.",
             "Screenshot script controller OTP & log pengiriman email",
             "Selesai (Verifikasi email OTP berfungsi)"),
            (5, "Front-End", "Desain Halaman Verifikasi OTP & Inisiasi Halaman Donasi",
             "Merancang tampilan form input 6-digit OTP responsif (verify-otp.blade.php) serta desain antarmuka awal halaman donasi publik (donasi.blade.php).",
             "Screenshot antarmuka verifikasi OTP & halaman donasi",
             "Selesai (UI OTP & Donasi terpasang)")
        ]
    },
    {
        "week_num": "2",
        "title": "Minggu ke : 2 | Hari : Senin – Jumat | Tanggal : 10 Agustus – 14 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Front-End", "Pembaruan UI Navbar, Lonceng Notifikasi & Hero Section",
             "Merancang ulang komponen navigation bar dengan icon lonceng notifikasi dinamis, penyesuaian menu aktif per role, dan desain hero section Beranda modern.",
             "Screenshot komponen navbar & hero section beranda",
             "Selesai (UI navigasi lebih interaktif)"),
            (2, "Back-End", "Refaktorisasi Sistem Hak Akses (Role & Permission)",
             "Refaktor tabel roles dan users untuk menetapkan 4 role resmi (Admin, Operator, Supir, Masyarakat), pembuatan gate authorization dan redirection login.",
             "Screenshot skema refaktor role di database & seeder",
             "Selesai (Hak akses role terstruktur)"),
            (3, "Front-End", "Pengembangan Landing Page Khusus Supir & Mitra Lembaga",
             "Membuat halaman informasi publik dan registrasi terpisah untuk supir ambulans (register-supir.blade.php) dan mitra lembaga (register-mitra.blade.php).",
             "Screenshot landing page supir dan portal mitra",
             "Selesai (Halaman supir & mitra siap)"),
            (4, "Back-End", "Integrasi Backend Payment Gateway Midtrans (Snap API)",
             "Pembuatan migrasi tabel donasis, konfigurasi environment key Midtrans, controller pembuatan Snap Token transaksi, dan webhook handler status pembayaran.",
             "Screenshot controller Midtrans & webhook endpoint",
             "Selesai (Payment gateway terintegrasi)"),
            (5, "Front-End", "Integrasi Modal Checkout Midtrans Snap & Tampilan Mitra",
             "Pemasangan script Midtrans Snap JS pada tombol bayar donasi, modal checkout multi-metode (QRIS, VA), dan pembuatan tampilan dashboard awal mitra.",
             "Screenshot popup Midtrans Snap checkout di browser",
             "Selesai (Checkout donasi berjalan lancar)")
        ]
    },
    {
        "week_num": "3",
        "title": "Minggu ke : 3 | Hari : Senin – Jumat | Tanggal : 17 Agustus – 21 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Front-End", "Pengembangan Mobile Console Supir & Bottom Navigation",
             "Membangun antarmuka dashboard supir (tugas-index.blade.php) yang dioptimasi untuk layar smartphone dengan mobile bottom navigation bar.",
             "Screenshot tampilan dashboard supir versi mobile smartphone",
             "Selesai (Console mobile supir responsif)"),
            (2, "Back-End", "Logika Triage Prioritas Darurat & Upload SIM/STNK",
             "Menambahkan field prioritas pemesanan (darurat/non-darurat) pada backend pemesanan, validasi upload file SIM & STNK supir, dan pagination data.",
             "Screenshot migration field prioritas & logic controller",
             "Selesai (Logika prioritas & berkas supir aktif)"),
            (3, "Back-End", "Setup REST API Token (Sanctum) & Logika Dispatching",
             "Konfigurasi Laravel Sanctum untuk endpoint API autentikasi dan pembuatan service logika dispatching penugasan pesanan ke supir terdekat.",
             "Screenshot route API Sanctum & PemesananService.php",
             "Selesai (API token & dispatching siap)"),
            (4, "Front-End", "Antarmuka Real-Time Tracking Ambulans & Rute Peta",
             "Membangun halaman live tracking pasien (order-tracking.blade.php) dengan visualisasi rute polyline jalan dan pergerakan marker ambulans realtime.",
             "Screenshot halaman live tracking pasien dan rute RS",
             "Selesai (Live tracking peta berfungsi)"),
            (5, "Back-End", "Kalkulasi Jarak OSRM, Estimasi Waktu (ETA) & Rating",
             "Integrasi API OSRM untuk kalkulasi jarak tempuh & perkiraan waktu tiba (ETA), serta migrasi tabel ratings untuk feedback kepuasan layanan.",
             "Screenshot kalkulasi ETA OSRM & tabel ratings",
             "Selesai (Kalkulasi rute & rating aktif)")
        ]
    },
    {
        "week_num": "4",
        "title": "Minggu ke : 4 | Hari : Senin – Jumat | Tanggal : 24 Agustus – 28 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Front-End", "Penyeragaman & Interaktivitas Peta Leaflet Seluruh Role",
             "Menyeragamkan tile layer OpenStreetMap di semua blade role (Admin, Operator, Mitra, Supir, Pasien), mengaktifkan zoom scroll mouse, dan merapikan pin marker 3D.",
             "Screenshot perbandingan tampilan peta seragam di semua role",
             "Selesai (Peta seragam, bersih, & interaktif)"),
            (2, "Back-End", "Validasi Alur Bisnis End-to-End & Sinkronisasi GPS",
             "Pengujian komprehensif alur status pemesanan (menunggu -> diproses -> menuju_lokasi -> membawa_pasien -> selesai) serta sinkronisasi koordinat GPS supir.",
             "Screenshot log status pemesanan & sinkronisasi GPS supir",
             "Selesai (Alur end-to-end terverifikasi)")
        ]
    }
]

# Styles
title_font = Font(name='Segoe UI', size=14, bold=True, color='0F2742')
subtitle_font = Font(name='Segoe UI', size=11, bold=True, color='1F497D')
header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
cell_font = Font(name='Segoe UI', size=10, color='1E293B')
status_font = Font(name='Segoe UI', size=10, bold=True, color='166534')

# Category fonts & fills
fe_font = Font(name='Segoe UI', size=10, bold=True, color='1E40AF') # Blue
fe_fill = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid')

be_font = Font(name='Segoe UI', size=10, bold=True, color='854D0E') # Amber/Brown
be_fill = PatternFill(start_color='FEFCE8', end_color='FEFCE8', fill_type='solid')

header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
status_fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')

thin_border_side = Side(border_style='thin', color='CBD5E1')
cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

header_border_side = Side(border_style='medium', color='1F497D')
header_border = Border(left=thin_border_side, right=thin_border_side, top=header_border_side, bottom=header_border_side)

for w in weeks_data:
    ws = wb.create_sheet(title=f"Minggu {w['week_num']}")
    ws.views.sheetView[0].showGridLines = True

    # Title
    ws.cell(row=1, column=1, value="Form Detail Pekerjaan").font = title_font
    ws.cell(row=2, column=1, value=w['title']).font = subtitle_font

    # Headers
    headers = ["No", "Kategori", "Pekerjaan/Tugas", "Uraian Pekerjaan", "Gambar", "Keterangan"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = header_border

    ws.row_dimensions[4].height = 28

    # Rows
    for row_idx, task in enumerate(w['tasks'], start=5):
        no, cat, name, desc, img, status = task
        
        c1 = ws.cell(row=row_idx, column=1, value=no)
        c1.alignment = Alignment(horizontal='center', vertical='center')
        c1.font = cell_font
        c1.border = cell_border

        c_cat = ws.cell(row=row_idx, column=2, value=cat)
        c_cat.alignment = Alignment(horizontal='center', vertical='center')
        if cat == "Front-End":
            c_cat.font = fe_font
            c_cat.fill = fe_fill
        else:
            c_cat.font = be_font
            c_cat.fill = be_fill
        c_cat.border = cell_border

        c2 = ws.cell(row=row_idx, column=3, value=name)
        c2.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        c2.font = Font(name='Segoe UI', size=10, bold=True, color='0F172A')
        c2.border = cell_border

        c3 = ws.cell(row=row_idx, column=4, value=desc)
        c3.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        c3.font = cell_font
        c3.border = cell_border

        c4 = ws.cell(row=row_idx, column=5, value=img)
        c4.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        c4.font = cell_font
        c4.border = cell_border

        c5 = ws.cell(row=row_idx, column=6, value=status)
        c5.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c5.font = status_font
        c5.fill = status_fill
        c5.border = cell_border

        ws.row_dimensions[row_idx].height = 55

    # Column widths
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 34
    ws.column_dimensions['D'].width = 58
    ws.column_dimensions['E'].width = 38
    ws.column_dimensions['F'].width = 24

# Create Summary Sheet: Front-End vs Back-End
ws_summary = wb.create_sheet(title="Ringkasan FE vs BE", index=0)
ws_summary.views.sheetView[0].showGridLines = True
ws_summary.cell(row=1, column=1, value="PEMBAGIAN TUGAS: FRONT-END VS BACK-END").font = title_font
ws_summary.cell(row=2, column=1, value="Proyek Sistem Informasi & Layanan Darurat Ambulans Siaga (Multi-Mitra)").font = subtitle_font

# Section 1: Front-End
ws_summary.cell(row=4, column=1, value="🎨 DAFTAR PEKERJAAN FRONT-END (UI/UX & CLIENT-SIDE)").font = Font(name='Segoe UI', size=11, bold=True, color='1E40AF')
fe_headers = ["No", "Modul / Fitur", "Teknologi", "Uraian Pekerjaan Front-End"]
for col_idx, h in enumerate(fe_headers, start=1):
    cell = ws_summary.cell(row=5, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = header_border

fe_rows = [
    (1, "Formulir Pemesanan Darurat", "Blade, Tailwind, Leaflet.js", "Membuat antarmuka input data pasien, pin marker penjemputan interaktif pada peta, dan modal verifikasi foto kondisi darurat."),
    (2, "Landing Page & Beranda", "Blade, Tailwind CSS, Alpine.js", "Merancang tampilan Hero section modern, alur layanan darurat, kartu keunggulan armada, dan layout responsif mobile."),
    (3, "Portal Supir & Mobile Nav", "Blade, Tailwind, Alpine.js", "Mendesain console dashboard supir dengan mobile bottom navigation bar yang nyaman digunakan di smartphone saat berkendara."),
    (4, "Live Tracking Peta Pasien", "Leaflet.js, Alpine.js, CSS", "Membuat visualisasi pergerakan posisi ambulans, garis rute OSRM ke rumah sakit, dan indikator status waktu kedatangan (ETA)."),
    (5, "Checkout Donasi Midtrans", "Midtrans Snap JS, Tailwind", "Integrasi antarmuka modal checkout pembayaran donasi multi-metode (QRIS, GoPay, Virtual Account Bank) dan tampilan kwitansi bukti transfer."),
    (6, "Dashboard Admin & Operator", "Blade, Tailwind, Leaflet.js", "Merancang antarmuka dispatch console, peta radar ambulans aktif, kartu statistik ringkasan, dan tabel riwayat order.")
]

for r_idx, r_data in enumerate(fe_rows, start=6):
    for c_idx, val in enumerate(r_data, start=1):
        cell = ws_summary.cell(row=r_idx, column=c_idx, value=val)
        cell.font = cell_font
        cell.border = cell_border
        if c_idx == 1:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    ws_summary.row_dimensions[r_idx].height = 36

# Section 2: Back-End
start_be_row = 14
ws_summary.cell(row=start_be_row, column=1, value="⚙️ DAFTAR PEKERJAAN BACK-END (SERVER-SIDE, DATABASE & API)").font = Font(name='Segoe UI', size=11, bold=True, color='854D0E')
be_headers = ["No", "Modul / Fitur", "Teknologi", "Uraian Pekerjaan Back-End"]
for col_idx, h in enumerate(be_headers, start=1):
    cell = ws_summary.cell(row=start_be_row+1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = PatternFill(start_color='B45309', end_color='B45309', fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = header_border

be_rows = [
    (1, "Arsitektur Database & Migrasi", "MySQL, Laravel Migration, Eloquent", "Merancang skema database (users, roles, supir, ambulans, pemesanan, donasis, mitras), relasi tabel, dan indeks performa query."),
    (2, "Sistem Autentikasi & OTP Email", "Laravel Auth, SMTP Mail, Session", "Membangun logika registrasi multi-role, pembuatan dan validasi kode OTP email acak, middleware otorisasi hak akses."),
    (3, "Payment Gateway & Webhook", "Midtrans Snap API, HTTP Client", "Menghubungkan API Midtrans dengan Server Key aman, membuat token Snap transaksi, serta menangani webhook callback perubahan status pembayaran."),
    (4, "Dispatching & Alur Pemesanan", "Livewire Component, Service Layer", "Mengembangkan algoritma penugasan armada ambulans terdekat ke pasien dan pengelolaan status alur darurat (menunggu -> diproses -> selesai)."),
    (5, "REST API GPS & Laravel Sanctum", "Laravel Sanctum, JSON API", "Membuat endpoint API token untuk pengiriman koordinat GPS driver secara berkala dan integrasi layanan routing OSRM."),
    (6, "Log Aktivitas & Pelaporan", "Laravel Eloquent, Event Listener", "Pencatatan riwayat aksi sistem (Audit Trail / LogAktivitas) serta kalkulasi rekapitulasi data statistik keuangan dan performa armada.")
]

for r_idx, r_data in enumerate(be_rows, start=start_be_row+2):
    for c_idx, val in enumerate(r_data, start=1):
        cell = ws_summary.cell(row=r_idx, column=c_idx, value=val)
        cell.font = cell_font
        cell.border = cell_border
        if c_idx == 1:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    ws_summary.row_dimensions[r_idx].height = 36

ws_summary.column_dimensions['A'].width = 6
ws_summary.column_dimensions['B'].width = 28
ws_summary.column_dimensions['C'].width = 32
ws_summary.column_dimensions['D'].width = 65

try:
    wb.save("laporan_mingguan_magang.xlsx")
    print("Saved as laporan_mingguan_magang.xlsx")
except Exception as e:
    print("Main file was open, saving to v2...")

wb.save("laporan_mingguan_magang_v2.xlsx")
print("Native Excel file saved successfully as laporan_mingguan_magang_v2.xlsx")
