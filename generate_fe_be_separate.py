import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
wb.remove(wb.active) # remove default sheet

# Data separated by Role & Week
weeks_fe = [
    {
        "week_num": "1",
        "sheet_title": "M1 - Front End",
        "title": "Form Detail Pekerjaan | Front-End Developer\nMinggu ke : 1 | Hari : Senin – Jumat | Tanggal : 3 Agustus – 7 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Slicing Master Layout & Setup Desain Sistem Tailwind",
             "Menyiapkan master layout (layouts/app.blade.php), konfigurasi warna tema Ambulance Siaga dengan Tailwind CSS, komponen header navbar dinamis, dan footer responsif.",
             "Screenshot master layout Blade & konfigurasi Tailwind",
             "Selesai (Master layout & desain sistem siap)"),
            (2, "Desain Formulir Pemesanan & Integrasi Peta Leaflet",
             "Membangun antarmuka formulir pemesanan darurat (order-create.blade.php), integrasi peta Leaflet.js dengan pin marker penjemputan draggable, dan tombol kontrol zoom.",
             "Screenshot UI form pemesanan & picker peta Leaflet",
             "Selesai (Formulir & peta penjemputan berfungsi)"),
            (3, "Desain Antarmuka Login, Registrasi & Upload Foto Darurat",
             "Membuat UI form login dan register multi-role, modal preview upload foto kondisi gawat darurat pasien, serta validasi visual error input form.",
             "Screenshot halaman login/register & preview upload foto",
             "Selesai (Antarmuka auth & upload foto aktif)"),
            (4, "Desain Halaman Verifikasi OTP & Landing Page Awal",
             "Mendesain form 6-digit input verifikasi kode OTP (verify-otp.blade.php) dengan styling countdown timer, serta layout awal halaman beranda publik.",
             "Screenshot UI verifikasi OTP & halaman beranda awal",
             "Selesai (UI OTP & Beranda awal terpasang)"),
            (5, "Desain Antarmuka Halaman Donasi Publik",
             "Membangun antarmuka halaman donasi (donasi.blade.php), kartu pilihan cepat nominal donasi (10k, 50k, 100k), dan form input data donatur.",
             "Screenshot antarmuka halaman donasi publik",
             "Selesai (Halaman donasi terpasang rapi)")
        ]
    },
    {
        "week_num": "2",
        "sheet_title": "M2 - Front End",
        "title": "Form Detail Pekerjaan | Front-End Developer\nMinggu ke : 2 | Hari : Senin – Jumat | Tanggal : 10 Agustus – 14 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Pembaruan UI Navbar & Lonceng Notifikasi Dinamis",
             "Merancang ulang komponen navigation bar dengan icon lonceng notifikasi dinamis (Alpine.js), menu aktif per role, dan penataan judul header.",
             "Screenshot navbar notifikasi & menu navigasi per role",
             "Selesai (Navbar notifikasi lebih interaktif)"),
            (2, "Desain Ulang Hero Section & Tampilan Beranda Modern",
             "Mengembangkan tampilan modern halaman Beranda (Hero section interaktif, badge status ambulans siaga 24 jam, kartu fitur, dan desain responsif mobile).",
             "Screenshot tampilan baru Landing Page Beranda",
             "Selesai (Landing page responsif & modern)"),
            (3, "Pengembangan Landing Page Portal Supir & Mitra Lembaga",
             "Membangun halaman informasi publik dan registrasi terpisah untuk supir ambulans (register-supir.blade.php) serta mitra lembaga (register-mitra.blade.php).",
             "Screenshot portal supir dan portal mitra lembaga",
             "Selesai (Portal supir & mitra siap digunakan)"),
            (4, "Desain Form Nominal Donasi & Komponen Payment Card",
             "Merancang antarmuka pilihan metode pembayaran donasi, form pesan/doa donatur, dan tombol bayar donasi interaktif.",
             "Screenshot form donasi & kartu pilihan pembayaran",
             "Selesai (Form pembayaran donasi siap)"),
            (5, "Integrasi Modal Checkout Midtrans Snap & Dashboard Mitra",
             "Pemasangan script Midtrans Snap JS pada tombol bayar, popup checkout pembayaran multi-metode (QRIS, VA), dan tampilan awal dashboard mitra.",
             "Screenshot modal Midtrans Snap checkout di browser",
             "Selesai (Checkout donasi berjalan lancar)")
        ]
    },
    {
        "week_num": "3",
        "sheet_title": "M3 - Front End",
        "title": "Form Detail Pekerjaan | Front-End Developer\nMinggu ke : 3 | Hari : Senin – Jumat | Tanggal : 17 Agustus – 21 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Desain Console Supir & Mobile Bottom Navigation Bar",
             "Membangun antarmuka dashboard supir (tugas-index.blade.php) yang dioptimasi untuk layar smartphone dengan mobile bottom navigation bar.",
             "Screenshot tampilan dashboard supir versi smartphone",
             "Selesai (Console mobile supir responsif)"),
            (2, "Desain Form Upload SIM/STNK & Badge Prioritas Order",
             "Membuat komponen form upload berkas foto SIM & STNK supir, preview berkas, serta visualisasi badge status prioritas order (Gawat Darurat / Normal).",
             "Screenshot form upload SIM/STNK & badge prioritas order",
             "Selesai (Komponen berkas & prioritas siap)"),
            (3, "Desain Antarmuka Dispatching Operator & Modal Penugasan",
             "Merancang UI konsol dispatch operator (order-index.blade.php), modal interaktif pemilihan supir terdekat, dan filter status panggilan.",
             "Screenshot konsol dispatching operator & modal assign",
             "Selesai (Dispatch console operator aktif)"),
            (4, "Desain Antarmuka Live Tracking GPS Pasien & Supir",
             "Membangun UI halaman tracking pasien (order-tracking.blade.php) dan supir (order-show.blade.php) dengan visualisasi polyline rute jalan dan animasi marker ambulans.",
             "Screenshot live tracking peta pasien & rute jalan RS",
             "Selesai (Live tracking peta berfungsi)"),
            (5, "Desain Form Rating Ulasan Bintang & Kwitansi Donasi",
             "Merancang modal rating bintang dan review feedback kepuasan layanan ambulans, serta halaman kwitansi digital bukti donasi online.",
             "Screenshot form rating bintang ulasan & kwitansi donasi",
             "Selesai (Form rating & kwitansi selesai)")
        ]
    },
    {
        "week_num": "4",
        "sheet_title": "M4 - Front End",
        "title": "Form Detail Pekerjaan | Front-End Developer\nMinggu ke : 4 | Hari : Senin – Jumat | Tanggal : 24 Agustus – 28 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Penyeragaman Visual Peta Leaflet & Interaktivitas Zoom",
             "Menyeragamkan tile layer OpenStreetMap di seluruh role (Admin, Operator, Mitra, Supir, Pasien), mengaktifkan zoom scroll mouse pada Admin, dan memperbarui pin marker 3D.",
             "Screenshot tampilan peta seragam di Admin, Operator, & Pasien",
             "Selesai (Peta seragam, bersih, & interaktif)"),
            (2, "Testing & Optimasi Responsivitas Antarmuka Seluruh Perangkat",
             "Melakukan pengujian antarmuka UI di berbagai ukuran layar (Desktop, Tablet, Mobile) dan merapikan komponen modal popup serta navigasi.",
             "Screenshot tampilan responsif mobile dan desktop",
             "Selesai (Antarmuka teroptimasi responsif)")
        ]
    }
]

weeks_be = [
    {
        "week_num": "1",
        "sheet_title": "M1 - Back End",
        "title": "Form Detail Pekerjaan | Back-End Developer\nMinggu ke : 1 | Hari : Senin – Jumat | Tanggal : 3 Agustus – 7 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Inisialisasi Proyek Laravel & Arsitektur Database Dasar",
             "Setup framework Laravel 11, konfigurasi database MySQL, pembuatan skema migrasi tabel dasar (users, roles, masyarakat, supir, ambulans, rumah_sakit), dan pemodelan relasi Eloquent.",
             "Screenshot struktur migration database & konfigurasi .env",
             "Selesai (Database & core Laravel siap)"),
            (2, "Logika Penyimpanan Data Pemesanan & Geocoding Koordinat",
             "Membuat Livewire Component OrderCreate.php, validasi input form pemesanan, penyimpanan koordinat latitude/longitude jemput, dan generator kode pemesanan otomatis.",
             "Screenshot class OrderCreate.php & tabel pemesanan",
             "Selesai (Backend create order berfungsi)"),
            (3, "Logika Autentikasi Multi-Role & Upload File Berkas",
             "Mengembangkan controller autentikasi login/register multi-role, middleware otorisasi hak akses, dan logic penyimpanan file foto kondisi gawat darurat ke local storage.",
             "Screenshot controller auth & logic upload file",
             "Selesai (Autentikasi & upload foto aktif)"),
            (4, "Sistem Generator OTP & Pengiriman Email Notifikasi",
             "Membuat migrasi tabel otp_verifications, logika pembuatan 6-digit OTP acak, verifikasi masa kedaluwarsa token, dan integrasi SMTP Laravel Mail.",
             "Screenshot script controller OTP & log email notifikasi",
             "Selesai (Verifikasi email OTP berfungsi)"),
            (5, "Logika Transaksi Donasi Awal & Sinkronisasi Repositori",
             "Membuat model Donasi, logic create transaksi donasi status pending, pengujian alur order darurat, serta menangani merge conflict branch git tim.",
             "Screenshot model Donasi.php & log git merge",
             "Selesai (Modul donasi aktif & repo sinkron)")
        ]
    },
    {
        "week_num": "2",
        "sheet_title": "M2 - Back End",
        "title": "Form Detail Pekerjaan | Back-End Developer\nMinggu ke : 2 | Hari : Senin – Jumat | Tanggal : 10 Agustus – 14 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Refaktorisasi Sistem Role & Permission Multi-Level",
             "Membuat migrasi refaktor 4 role resmi (Admin, Operator, Supir, Masyarakat), seeder data role, dan logic role-based redirection saat pengguna login.",
             "Screenshot skema refaktor role di database & seeder",
             "Selesai (Hak akses role terstruktur)"),
            (2, "Penyempurnaan Query Data Agregasi Beranda",
             "Membuat query statistik dinamis pada halaman beranda (total armada siaga, total panggilan tertangani, dan daftar RS mitra rujukan).",
             "Screenshot controller query agregasi data beranda",
             "Selesai (Query statistik beranda aktif)"),
            (3, "Backend Pendaftaran Supir & Mitra Lembaga",
             "Membuat controller RegisteredSupirController.php dan RegisteredMitraController.php, validasi form registrasi data lembaga, dan relasi mitra ke armada/driver.",
             "Screenshot controller registrasi supir & mitra",
             "Selesai (Backend supir & mitra siap)"),
            (4, "Perancangan Skema Database Donasi & Konfigurasi API Midtrans",
             "Membuat migrasi tabel donasis, konfigurasi environment key Midtrans (Client Key & Server Key) pada .env dan config/services.php.",
             "Screenshot migrasi tabel donasis & konfigurasi Midtrans",
             "Selesai (Database & konfigurasi API siap)"),
            (5, "Implementasi Midtrans Snap API & Webhook Handler",
             "Membangun endpoint pembuatan Snap Token transaksi donasi, migrasi tabel mitras, serta webhook controller untuk auto-update status pembayaran donasi (settlement/expire).",
             "Screenshot webhook controller & log transaksi Midtrans",
             "Selesai (Payment gateway terintegrasi)")
        ]
    },
    {
        "week_num": "3",
        "sheet_title": "M3 - Back End",
        "title": "Form Detail Pekerjaan | Back-End Developer\nMinggu ke : 3 | Hari : Senin – Jumat | Tanggal : 17 Agustus – 21 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Logika Penerimaan Tugas Supir & Status Perjalanan",
             "Membuat logika penerimaan/penolakan order oleh supir, update status perjalanan (menuju_lokasi, membawa_pasien, selesai), dan event listener perubahan status.",
             "Screenshot Livewire Supir OrderShow.php & status alur",
             "Selesai (Logika tugas supir berfungsi)"),
            (2, "Backend Prioritas Darurat, Verifikasi SIM/STNK & Pagination",
             "Menambahkan migrasi kolom prioritas pada tabel pemesanan, migrasi kolom foto_sim & foto_stnk pada tabel supir, dan optimalisasi pagination data order.",
             "Screenshot file migrasi prioritas & validasi berkas",
             "Selesai (Logika prioritas & berkas supir aktif)"),
            (3, "Setup REST API Token (Sanctum) & Logika Dispatching",
             "Konfigurasi Laravel Sanctum (personal_access_tokens), pembuatan service PemesananService.php untuk algoritma penugasan armada terdekat ke pasien.",
             "Screenshot PemesananService.php & route Sanctum API",
             "Selesai (API token & dispatching siap)"),
            (4, "Integrasi OSRM (Routing API), Estimasi ETA & Polling GPS",
             "Menghubungkan API OSRM untuk kalkulasi jarak kilometer dan estimasi waktu tempuh (ETA), serta logika polling sinkronisasi koordinat GPS armada.",
             "Screenshot fungsi kalkulasi rute OSRM & polling Livewire",
             "Selesai (Kalkulasi rute & live GPS aktif)"),
            (5, "Backend Modul Rating Feedback & Kwitansi Transaksi",
             "Membuat migrasi penyempurnaan tabel ratings, logika kalkulasi rata-rata kepuasan supir, dan query detail transaksi donasi untuk kwitansi digital.",
             "Screenshot migrasi rating & controller kwitansi donasi",
             "Selesai (Backend rating & kwitansi selesai)")
        ]
    },
    {
        "week_num": "4",
        "sheet_title": "M4 - Back End",
        "title": "Form Detail Pekerjaan | Back-End Developer\nMinggu ke : 4 | Hari : Senin – Jumat | Tanggal : 24 Agustus – 28 Agustus | Tahun : 2026",
        "tasks": [
            (1, "Optimasi Sinkronisasi Realtime Marker Status Armada",
             "Memperbaiki logic pemetaan status operasional supir di Dashboard Admin (Emergency, On Duty, Available, Offline) dan sinkronisasi koordinat terakhir supir.",
             "Screenshot query status supir di Admin/Dashboard.php",
             "Selesai (Pemetaan status armada presisi)"),
            (2, "Pengujian Alur Bisnis End-to-End & Audit Trail Log Aktivitas",
             "Melakukan testing komprehensif alur pemesanan darurat (pasien -> operator -> supir -> selesai) serta verifikasi pencatatan riwayat aksi pada tabel log_aktivitas.",
             "Screenshot log status pemesanan & tabel log_aktivitas",
             "Selesai (Alur end-to-end terverifikasi)")
        ]
    }
]

# Styles
title_font = Font(name='Segoe UI', size=13, bold=True, color='0F2742')
subtitle_font = Font(name='Segoe UI', size=10, bold=True, color='1F497D')
header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
cell_font = Font(name='Segoe UI', size=10, color='1E293B')
status_font = Font(name='Segoe UI', size=10, bold=True, color='166534')

# Palette
fe_header_fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid') # Deep Blue
be_header_fill = PatternFill(start_color='B45309', end_color='B45309', fill_type='solid') # Amber Brown
status_fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')

thin_border_side = Side(border_style='thin', color='CBD5E1')
cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

def populate_sheet(ws, item, is_fe=True):
    ws.views.sheetView[0].showGridLines = True
    
    # Title Block
    lines = item['title'].split('\n')
    ws.cell(row=1, column=1, value=lines[0]).font = Font(name='Segoe UI', size=13, bold=True, color='1E40AF' if is_fe else 'B45309')
    ws.cell(row=2, column=1, value=lines[1]).font = subtitle_font

    # Headers
    headers = ["No", "Pekerjaan/Tugas", "Uraian Pekerjaan", "Gambar", "Keterangan"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = fe_header_fill if is_fe else be_header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = cell_border

    ws.row_dimensions[4].height = 26

    # Rows
    for row_idx, task in enumerate(item['tasks'], start=5):
        no, name, desc, img, status = task
        
        c1 = ws.cell(row=row_idx, column=1, value=no)
        c1.alignment = Alignment(horizontal='center', vertical='center')
        c1.font = cell_font
        c1.border = cell_border

        c2 = ws.cell(row=row_idx, column=2, value=name)
        c2.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        c2.font = Font(name='Segoe UI', size=10, bold=True, color='0F172A')
        c2.border = cell_border

        c3 = ws.cell(row=row_idx, column=3, value=desc)
        c3.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        c3.font = cell_font
        c3.border = cell_border

        c4 = ws.cell(row=row_idx, column=4, value=img)
        c4.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        c4.font = cell_font
        c4.border = cell_border

        c5 = ws.cell(row=row_idx, column=5, value=status)
        c5.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c5.font = status_font
        c5.fill = status_fill
        c5.border = cell_border

        ws.row_dimensions[row_idx].height = 52

    # Column widths
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 60
    ws.column_dimensions['D'].width = 38
    ws.column_dimensions['E'].width = 24

# Create paired sheets in chronological order: M1 FE, M1 BE, M2 FE, M2 BE, M3 FE, M3 BE, M4 FE, M4 BE
for w_idx in range(4):
    # Front End Sheet
    ws_fe = wb.create_sheet(title=weeks_fe[w_idx]['sheet_title'])
    populate_sheet(ws_fe, weeks_fe[w_idx], is_fe=True)

    # Back End Sheet
    ws_be = wb.create_sheet(title=weeks_be[w_idx]['sheet_title'])
    populate_sheet(ws_be, weeks_be[w_idx], is_fe=False)

output_file = "laporan_mingguan_magang_fe_be_terpisah.xlsx"
wb.save(output_file)
print(f"Successfully generated: {output_file}")
